from pathlib import Path

from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument


def load_document(file_path):
    """
    Load TXT, PDF, or DOCX documents.
    Returns a list of LangChain Document objects.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = file_path.suffix.lower()

    # -------------------------
    # TXT
    # -------------------------

    if extension == ".txt":

        text = file_path.read_text(
            encoding="utf-8-sig"
        )

        return [
            Document(
                page_content=text,
                metadata={
                    "source": file_path.name,
                    "file_type": "txt"
                }
            )
        ]

    # -------------------------
    # PDF
    # -------------------------

    elif extension == ".pdf":

        reader = PdfReader(str(file_path))

        documents = []

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            text = page.extract_text() or ""

            if text.strip():

                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": file_path.name,
                            "file_type": "pdf",
                            "page": page_number
                        }
                    )
                )

        return documents

    # -------------------------
    # DOCX
    # -------------------------

    elif extension == ".docx":

        docx = DocxDocument(str(file_path))

        text = "\n".join(
            paragraph.text
            for paragraph in docx.paragraphs
            if paragraph.text.strip()
        )

        return [
            Document(
                page_content=text,
                metadata={
                    "source": file_path.name,
                    "file_type": "docx"
                }
            )
        ]

    else:

        raise ValueError(
            "Unsupported file type. "
            "Use PDF, DOCX, or TXT."
        )